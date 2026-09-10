"""Tests for the PPST output workbook builder on a tiny synthetic run (no DB)."""

from __future__ import annotations

import openpyxl
import pytest

from cfde_atlas_etl.pubsearch.export_output import (
    QUERY_SUMMARY_HEADER,
    TABULAR_HEADER,
    build_workbook,
)

RUN_ID = "cfde-abc1234-20260910T120000"
RUN = {
    "run_id": RUN_ID,
    "program": "cfde",
    "program_yaml_sha": "abc1234",
    "date_window_start": 2020,
    "date_window_end": 2026,
    "started_at": "2026-09-10 12:00:00+00:00",
}


def _q(no: int, tier: str, cluster: str) -> dict:
    return {
        "run_id": RUN_ID,
        "query_no": no,
        "program": "cfde",
        "impact_category": tier,
        "query_cluster": cluster,
        "search_terms": f"term{no}",
        "search_field": "",
        "epmc_query": f"({cluster}:term{no})",
        "hit_count": 10 * no,
    }


QUERIES = [
    _q(1, "Awardee", "grant"),
    _q(2, "User", "resource_url"),
    _q(3, "Broader.Influence", "cites"),
]


def _hit(no: int, source: str, epmc_id: str, pmid: str | None, year: int) -> dict:
    return {
        "run_id": RUN_ID,
        "query_no": no,
        "epmc_id": epmc_id,
        "source": source,
        "pmid": pmid,
        "pmcid": None,
        "pub_year": year,
        "is_open_access": True,
        "title": f"Paper {epmc_id}",
    }


# p1: awardee+user (q1,q2); p2: user only, overridden to Awardee; p3: broader; PPR1: preprint via q2.
HITS = [
    _hit(1, "MED", "1001", "1001", 2021),
    _hit(2, "MED", "1001", "1001", 2021),
    _hit(2, "MED", "1002", "1002", 2022),
    _hit(3, "MED", "1003", "1003", 2023),
    _hit(2, "PPR", "PPR1", None, 2024),
]


def _m(key: str, pmid: str | None, year: int, nos: list[int], clusters: list[str], **kw) -> dict:
    a, u, b = 1 in nos, 2 in nos, 3 in nos
    computed = "Awardee" if a else "User" if u else "Broader.Influence"
    override = kw.get("override_tier")
    return {
        "run_id": RUN_ID,
        "program": "cfde",
        "paper_key": key,
        "pmid": pmid,
        "pmcid": None,
        "source": "PPR" if pmid is None else "MED",
        "is_preprint": pmid is None,
        "pub_year": year,
        "title": f"Paper {key}",
        "is_open_access": True,
        "awardee": a,
        "user_tier": u,
        "broader_influence": b,
        "computed_assignment": computed,
        "override_tier": override,
        "final_assignment": override or computed,
        "n_queries": len(nos),
        "clusters": clusters,
        "query_nos": nos,
        "evidence": kw.get("evidence"),
    }


MATRIX = [
    _m(
        "1001",
        "1001",
        2021,
        [1, 2],
        ["grant", "resource_url"],
        evidence="Data from the CFDE portal.",
    ),
    _m("1002", "1002", 2022, [2], ["resource_url"], override_tier="Awardee"),
    _m("1003", "1003", 2023, [3], ["cites"]),
    _m("PPR:PPR1", None, 2024, [2], ["resource_url"]),
]


@pytest.fixture
def wb() -> openpyxl.Workbook:
    return build_workbook(RUN, QUERIES, MATRIX, HITS)


def _rows(ws) -> list[list]:
    """Rows with trailing blank cells stripped (iter_rows pads to the sheet width)."""
    out = []
    for r in ws.iter_rows(values_only=True):
        r = list(r)
        while r and r[-1] is None:
            r.pop()
        out.append(r)
    return out


def test_sheet_names_and_order(wb):
    assert wb.sheetnames == [
        "Summary",
        "Data_Tabular",
        "Data_Matrixed_Query",
        "Data_Matrixed_Cluster",
    ]


def test_headers(wb):
    tier = ["Awardee", "User", "Broader_Influence", "Final_Assignment"]
    extra = ["Source", "Is_Preprint", "Override_Tier", "Evidence", "Run_ID"]
    assert _rows(wb["Data_Tabular"])[0] == list(TABULAR_HEADER)
    assert _rows(wb["Data_Matrixed_Query"])[0] == ["pmid", "pubYear", "1", "2", "3", *tier, *extra]
    assert _rows(wb["Data_Matrixed_Cluster"])[0] == [
        "pmid", "pubYear", "grant", "resource_url", "cites", *tier, *extra
    ]  # fmt: skip
    summary = _rows(wb["Summary"])
    assert summary[9][:3] == ["Impact_Category", "Query_clusters", "Total_PMIDs"]
    assert summary[15][:4] == ["Query_Cluster", "Impact_Category", "Query_Numbers", "Total_PMIDs"]
    assert summary[21] == list(QUERY_SUMMARY_HEADER)


def test_matrix_query_header_prefix(wb):
    assert _rows(wb["Data_Matrixed_Query"])[0][:5] == ["pmid", "pubYear", "1", "2", "3"]


def test_matrix_x_placement_and_final_assignment(wb):
    rows = {r[0]: r for r in _rows(wb["Data_Matrixed_Query"])[1:]}
    assert set(rows) == {"1001", "1002", "1003", "PPR:PPR1"}
    assert rows["1001"][2:9] == ["x", "x", None, "x", "x", None, "Awardee"]
    assert rows["1002"][2:9] == [None, "x", None, None, "x", None, "Awardee"]  # override wins
    assert rows["1002"][11] == "Awardee"  # Override_Tier
    assert rows["1003"][2:9] == [None, None, "x", None, None, "x", "Broader.Influence"]
    assert rows["PPR:PPR1"][2:9] == [None, "x", None, None, "x", None, "User"]
    assert rows["PPR:PPR1"][9:11] == ["PPR", True]
    assert rows["1001"][12] == "Data from the CFDE portal."
    assert all(r[13] == RUN_ID for r in rows.values())

    crows = {r[0]: r for r in _rows(wb["Data_Matrixed_Cluster"])[1:]}
    assert crows["1001"][2:5] == ["x", "x", None]
    assert crows["1003"][2:5] == [None, None, "x"]


def test_summary_counts(wb):
    s = _rows(wb["Summary"])
    assert s[0][0] == "Date & Time:"
    assert s[2] == ["Period Covered:", "2020-2026"]
    assert s[3] == ["Analyst Name:", f"cfde-atlas pubsearch {RUN_ID}"]
    assert s[4:7] == [
        ["Run_ID:", RUN_ID],
        ["Program_YAML_SHA:", "abc1234"],
        ["Preprints_Included:", True],
    ]
    # tier rows: pmid counts by final_assignment (preprint has no pmid)
    assert s[10:13] == [
        ["Awardee", "grant", 2],
        ["User", "resource_url", 0],
        ["Broader.Influence", "cites", 1],
    ]
    assert s[16:19] == [
        ["grant", "Awardee", "1", 1],
        ["resource_url", "User", "2", 2],
        ["cites", "Broader.Influence", "3", 1],
    ]


def test_query_summary_total_pmids_excludes_preprints(wb):
    s = _rows(wb["Summary"])
    q = {r[0]: r for r in s[22:25]}
    assert q[2][:6] == [2, "resource_url", "resource_url", "User", 20, 2]  # PPR1 not counted
    assert q[1][6] == "(grant:term1)"


def test_data_tabular(wb):
    rows = _rows(wb["Data_Tabular"])[1:]
    assert len(rows) == 5
    assert rows[0][:3] == ["1001", "MED", "1001"]
    assert rows[0][7:11] == [1, "grant", "Awardee", "(grant:term1)"]
    assert rows[0][11] == "Data from the CFDE portal."
